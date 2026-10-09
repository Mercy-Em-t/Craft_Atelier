import os
import re
try:
    from docx import Document
    from docx.shared import Pt, RGBColor
    
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)
    
    files = ['README.md', 'SOP_WAREHOUSE.md', 'ARCHITECTURE.md']
    
    for idx, filename in enumerate(files):
        if idx > 0:
            doc.add_page_break()
            
        if not os.path.exists(filename):
            continue
            
        with open(filename, 'r', encoding='utf-8') as f:
            lines = f.readlines()
            
        in_code_block = False
        code_text = ""
        
        for line in lines:
            line = line.replace('\n', '')
            stripped = line.strip()
            
            if stripped.startswith('```'):
                if in_code_block:
                    p = doc.add_paragraph(code_text)
                    # p.style = 'Macro Text' or similar, we'll just use normal with Courier New
                    for run in p.runs:
                        run.font.name = 'Courier New'
                    in_code_block = False
                    code_text = ""
                else:
                    in_code_block = True
                continue
                
            if in_code_block:
                code_text += line + "\n"
                continue
                
            if stripped.startswith('# '):
                doc.add_heading(stripped[2:].replace('*', ''), level=1)
            elif stripped.startswith('## '):
                doc.add_heading(stripped[3:].replace('*', ''), level=2)
            elif stripped.startswith('### '):
                doc.add_heading(stripped[4:].replace('*', ''), level=3)
            elif stripped.startswith('* ') or stripped.startswith('- '):
                doc.add_paragraph(stripped[2:], style='List Bullet')
            elif re.match(r'^\d+\.\s', stripped):
                doc.add_paragraph(re.sub(r'^\d+\.\s', '', stripped), style='List Number')
            elif stripped == '':
                continue
            else:
                doc.add_paragraph(stripped)
    
    doc.save('Craft_Atelier_Documentation_Final.docx')
    print("Successfully compiled to Word document.")
except Exception as e:
    print(f"Error: {e}")
